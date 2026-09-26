"""
Build Excel/Healthcare_Claims_Analysis.xlsx from the clean claims data.

Every summary number in the workbook is a live Excel formula (COUNTIFS / SUMIFS /
AVERAGEIFS) that points at the Claims sheet, so a reviewer can click any cell and
see how it was calculated.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Excel" / "Healthcare_Claims_Analysis.xlsx"
OUT.parent.mkdir(exist_ok=True)

BLUE, ORANGE, AQUA, DARK = "2A78D6", "EB6834", "1BAF7A", "0D366B"
HEADER = PatternFill("solid", fgColor=DARK)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=18, color="0B0B0B")
SUB_FONT = Font(size=10, color="52514E")
THIN = Side(style="thin", color="E1E0D9")
MONEY, PCT, INT = '"$"#,##0', "0.0%", "#,##0"

# ---------------------------------------------------------------- data
claims = pd.read_csv(ROOT / "Data" / "clean" / "claims_clean.csv", dtype={"cpt_code": str})
providers = pd.read_csv(ROOT / "Data" / "clean" / "providers_clean.csv")
claims = claims.merge(providers[["provider_id", "provider_name"]], on="provider_id")
# numeric month key (202401) keeps COUNTIFS criteria unambiguous; text strings like "2024-01" can be read as dates
claims["month_key"] = claims["service_month"].str.replace("-", "").astype(int)

COLS = ["claim_id", "month_key", "service_date", "plan_type", "provider_id", "provider_name", "specialty",
        "network_status", "claim_type", "diagnosis_category", "billed_amount", "allowed_amount",
        "paid_amount", "claim_status", "denial_reason", "days_to_process", "age_group"]
data = claims[COLS].sort_values(["service_date", "claim_id"])
N = len(data)
LAST = N + 1
col = {name: get_column_letter(i + 1) for i, name in enumerate(COLS)}


def rng(name):
    """Absolute range for one Claims column, e.g. Claims!$N$2:$N$50001."""
    c = col[name]
    return f"Claims!${c}$2:${c}${LAST}"


def header_row(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=label)
        cell.fill, cell.font = HEADER, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, mapping):
    for letter, w in mapping.items():
        ws.column_dimensions[letter].width = w


def title(ws, text, sub):
    ws["A1"], ws["A2"] = text, sub
    ws["A1"].font, ws["A2"].font = TITLE_FONT, SUB_FONT


wb = Workbook()

# ---------------------------------------------------------------- Claims (raw data table)
ws_data = wb.active
ws_data.title = "Claims"
ws_data.append(COLS)
for row in data.itertuples(index=False):
    ws_data.append(list(row))
table = Table(displayName="tblClaims", ref=f"A1:{get_column_letter(len(COLS))}{LAST}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_data.add_table(table)
ws_data.freeze_panes = "B2"
for name in ["billed_amount", "allowed_amount", "paid_amount"]:
    for cell in ws_data[col[name]][1:]:
        cell.number_format = '"$"#,##0.00'
widths(ws_data, {get_column_letter(i + 1): max(12, len(c) + 3) for i, c in enumerate(COLS)})
widths(ws_data, {col["provider_name"]: 28, col["diagnosis_category"]: 20, col["denial_reason"]: 28})

# ---------------------------------------------------------------- Monthly
ws_m = wb.create_sheet("Monthly")
title(ws_m, "Monthly Trend", "Formulas: COUNTIFS / SUMIFS against the Claims sheet. Denial rate excludes pending claims.")
header_row(ws_m, 4, ["Month key", "Month", "Claims", "Denied", "Pending", "Denial rate", "Billed", "Paid", "Paid YTD"])
months = sorted(claims["month_key"].unique())
for i, key in enumerate(months):
    r = 5 + i
    ws_m[f"A{r}"] = int(key)
    ws_m[f"B{r}"] = pd.Timestamp(year=key // 100, month=key % 100, day=1).strftime("%b %Y")
    ws_m[f"C{r}"] = f"=COUNTIFS({rng('month_key')},A{r})"
    ws_m[f"D{r}"] = f'=COUNTIFS({rng("month_key")},A{r},{rng("claim_status")},"Denied")'
    ws_m[f"E{r}"] = f'=COUNTIFS({rng("month_key")},A{r},{rng("claim_status")},"Pending")'
    ws_m[f"F{r}"] = f"=IFERROR(D{r}/(C{r}-E{r}),0)"
    ws_m[f"G{r}"] = f"=SUMIFS({rng('billed_amount')},{rng('month_key')},A{r})"
    ws_m[f"H{r}"] = f"=SUMIFS({rng('paid_amount')},{rng('month_key')},A{r})"
    # year-to-date running total resets each January
    ws_m[f"I{r}"] = f"=IF(MOD(A{r},100)=1,H{r},I{r - 1}+H{r})" if i else f"=H{r}"
    for c, fmt in zip("CDEFGHI", [INT, INT, INT, PCT, MONEY, MONEY, MONEY]):
        ws_m[f"{c}{r}"].number_format = fmt
m_last = 4 + len(months)
widths(ws_m, {"A": 11, "B": 11, "C": 10, "D": 10, "E": 10, "F": 12, "G": 15, "H": 15, "I": 15})
ws_m.freeze_panes = "A5"

# ---------------------------------------------------------------- Denial reasons
ws_r = wb.create_sheet("Denial Reasons")
title(ws_r, "Denial Reasons", "Which reasons cost the most, and how they changed from 2024 to 2025.")
header_row(ws_r, 4, ["Denial reason", "Denied claims", "Denied billed $", "Share of denied $",
                     "Denials 2024", "Denials 2025", "YoY change"])
reasons = (claims[claims["claim_status"] == "Denied"].groupby("denial_reason")["billed_amount"]
           .sum().sort_values(ascending=False).index.tolist())
for i, reason in enumerate(reasons):
    r = 5 + i
    ws_r[f"A{r}"] = reason
    ws_r[f"B{r}"] = f'=COUNTIFS({rng("denial_reason")},A{r},{rng("claim_status")},"Denied")'
    ws_r[f"C{r}"] = f'=SUMIFS({rng("billed_amount")},{rng("denial_reason")},A{r},{rng("claim_status")},"Denied")'
    ws_r[f"D{r}"] = f"=C{r}/SUM($C$5:$C${4 + len(reasons)})"
    ws_r[f"E{r}"] = f'=COUNTIFS({rng("denial_reason")},A{r},{rng("claim_status")},"Denied",{rng("month_key")},">=202401",{rng("month_key")},"<=202412")'
    ws_r[f"F{r}"] = f'=COUNTIFS({rng("denial_reason")},A{r},{rng("claim_status")},"Denied",{rng("month_key")},">=202501",{rng("month_key")},"<=202512")'
    ws_r[f"G{r}"] = f"=IFERROR(F{r}/E{r}-1,0)"
    for c, fmt in zip("BCDEFG", [INT, MONEY, PCT, INT, INT, PCT]):
        ws_r[f"{c}{r}"].number_format = fmt
r_last = 4 + len(reasons)
ws_r.conditional_formatting.add(f"C5:C{r_last}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                              color=BLUE, showValue=True))
ws_r.conditional_formatting.add(f"G5:G{r_last}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                                 end_type="max", end_color="F4B183"))
widths(ws_r, {"A": 30, "B": 14, "C": 18, "D": 16, "E": 13, "F": 13, "G": 12})

# ---------------------------------------------------------------- Plan x Network
ws_p = wb.create_sheet("Plan x Network")
title(ws_p, "Denial Rate by Plan and Network", "Denied / decided claims. Out-of-network claims are denied about twice as often.")
header_row(ws_p, 4, ["Plan type", "In-Network", "Out-of-Network", "Gap (pts)"])
for i, plan in enumerate(["Commercial", "Medicare", "Medicaid"]):
    r = 5 + i
    ws_p[f"A{r}"] = plan
    for c, net in [("B", "In-Network"), ("C", "Out-of-Network")]:
        denied = f'COUNTIFS({rng("plan_type")},$A{r},{rng("network_status")},"{net}",{rng("claim_status")},"Denied")'
        decided = f'COUNTIFS({rng("plan_type")},$A{r},{rng("network_status")},"{net}",{rng("claim_status")},"<>Pending")'
        ws_p[f"{c}{r}"] = f"={denied}/{decided}"
        ws_p[f"{c}{r}"].number_format = PCT
    ws_p[f"D{r}"] = f"=(C{r}-B{r})*100"
    ws_p[f"D{r}"].number_format = "0.0"
ws_p.conditional_formatting.add("B5:C7", ColorScaleRule(start_type="num", start_value=0.05, start_color="FFFFFF",
                                                         end_type="num", end_value=0.25, end_color="EB6834"))
widths(ws_p, {"A": 16, "B": 14, "C": 16, "D": 11})

# ---------------------------------------------------------------- Provider scorecard
ws_s = wb.create_sheet("Provider Scorecard")
title(ws_s, "Provider Scorecard", "Filter or sort any column. Peer rate = providers with the same claim type and network status.")
header_row(ws_s, 4, ["Provider ID", "Provider", "Specialty", "Network", "Claim type", "Decided claims",
                     "Denied", "Denial rate", "Peer rate", "Excess (pts)", "Denied billed $", "Paid $"])
decided = claims[claims["claim_status"] != "Pending"]
peer = decided.groupby(["claim_type", "network_status"])["is_denied"].transform("mean")
prov = (decided.assign(excess=decided["is_denied"] - peer)
        .groupby(["provider_id", "provider_name", "specialty", "network_status", "claim_type"])["excess"].mean()
        .sort_values(ascending=False)       # worst offenders first; the formulas recompute the values
        .reset_index()[["provider_id", "provider_name", "specialty", "network_status", "claim_type"]])
s_last = 4 + len(prov)
for i, p in enumerate(prov.itertuples(index=False)):
    r = 5 + i
    for c, v in zip("ABCDE", p):
        ws_s[f"{c}{r}"] = v
    ws_s[f"F{r}"] = f'=COUNTIFS({rng("provider_id")},A{r},{rng("claim_status")},"<>Pending")'
    ws_s[f"G{r}"] = f'=COUNTIFS({rng("provider_id")},A{r},{rng("claim_status")},"Denied")'
    ws_s[f"H{r}"] = f"=IFERROR(G{r}/F{r},0)"
    ws_s[f"I{r}"] = f"=SUMIFS($G$5:$G${s_last},$D$5:$D${s_last},D{r},$E$5:$E${s_last},E{r})/SUMIFS($F$5:$F${s_last},$D$5:$D${s_last},D{r},$E$5:$E${s_last},E{r})"
    ws_s[f"J{r}"] = f"=(H{r}-I{r})*100"
    ws_s[f"K{r}"] = f'=SUMIFS({rng("billed_amount")},{rng("provider_id")},A{r},{rng("claim_status")},"Denied")'
    ws_s[f"L{r}"] = f"=SUMIFS({rng('paid_amount')},{rng('provider_id')},A{r})"
    for c, fmt in zip("FGHIJKL", [INT, INT, PCT, PCT, "0.0", MONEY, MONEY]):
        ws_s[f"{c}{r}"].number_format = fmt
ws_s.auto_filter.ref = f"A4:L{s_last}"
ws_s.conditional_formatting.add(f"J5:J{s_last}", ColorScaleRule(
    start_type="num", start_value=-10, start_color="2A78D6", mid_type="num", mid_value=0, mid_color="FFFFFF",
    end_type="num", end_value=20, end_color="EB6834"))
ws_s.freeze_panes = "C5"
widths(ws_s, {"A": 11, "B": 30, "C": 20, "D": 15, "E": 13, "F": 10, "G": 9, "H": 11, "I": 10, "J": 11, "K": 15, "L": 14})

# ---------------------------------------------------------------- Dashboard
ws_d = wb.create_sheet("Dashboard", 0)
ws_d.sheet_view.showGridLines = False
title(ws_d, "Healthcare Claims Dashboard, 2024-2025",
      "Where is claim revenue being lost? All figures are live formulas over the Claims sheet (synthetic data).")
kpis = [
    ("Total claims",        f"=COUNTA({rng('claim_id')})", INT),
    ("Total billed",        f"=SUM({rng('billed_amount')})", MONEY),
    ("Total paid",          f"=SUM({rng('paid_amount')})", MONEY),
    ("Denial rate",         f'=COUNTIFS({rng("claim_status")},"Denied")/COUNTIFS({rng("claim_status")},"<>Pending")', PCT),
    ("Denied billed $",     f'=SUMIFS({rng("billed_amount")},{rng("claim_status")},"Denied")', MONEY),
    ("Avg days to process", f"=AVERAGE({rng('days_to_process')})", "0.0"),
]
for i, (label, formula, fmt) in enumerate(kpis):
    c = get_column_letter(1 + i * 2)
    ws_d.merge_cells(f"{c}4:{get_column_letter(2 + i * 2)}4")
    ws_d.merge_cells(f"{c}5:{get_column_letter(2 + i * 2)}5")
    ws_d[f"{c}4"], ws_d[f"{c}5"] = label, formula
    ws_d[f"{c}4"].font = Font(size=10, color="52514E")
    ws_d[f"{c}5"].font = Font(size=18, bold=True, color="0B0B0B")
    ws_d[f"{c}5"].number_format = fmt
    for rr in (4, 5):
        ws_d[f"{c}{rr}"].alignment = Alignment(horizontal="left", indent=1)
    for cc in range(1 + i * 2, 3 + i * 2):
        ws_d.cell(row=4, column=cc).border = Border(top=Side(style="medium", color=BLUE))
for c in range(1, 13):
    ws_d.column_dimensions[get_column_letter(c)].width = 13
ws_d.row_dimensions[5].height = 32
ws_d.page_setup.orientation = "landscape"
ws_d.sheet_properties.pageSetUpPr.fitToPage = True
ws_d.page_setup.fitToWidth = ws_d.page_setup.fitToHeight = 1


def style_chart(ch, title_text, y_title=None):
    ch.title = title_text
    ch.title.overlay = False
    ch.style = 2
    ch.legend = None
    ch.height, ch.width = 7.5, 15
    if y_title:
        ch.y_axis.title = y_title
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E1E0D9"))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    return ch


# monthly claims line
line = LineChart()
line.add_data(Reference(ws_m, min_col=3, min_row=4, max_row=m_last), titles_from_data=True)
line.set_categories(Reference(ws_m, min_col=2, min_row=5, max_row=m_last))
style_chart(line, "Monthly claim volume")
line.series[0].graphicalProperties.line.solidFill = BLUE
line.series[0].graphicalProperties.line.width = 28000
line.series[0].smooth = False
ws_d.add_chart(line, "A8")

# monthly denial rate line (separate chart, never a second axis)
rate = LineChart()
rate.add_data(Reference(ws_m, min_col=6, min_row=4, max_row=m_last - 1), titles_from_data=True)
rate.set_categories(Reference(ws_m, min_col=2, min_row=5, max_row=m_last - 1))
style_chart(rate, "Monthly denial rate (Dec 2025 excluded: mostly pending)")
rate.y_axis.number_format = "0%"
rate.series[0].graphicalProperties.line.solidFill = ORANGE
rate.series[0].graphicalProperties.line.width = 28000
rate.series[0].smooth = False
ws_d.add_chart(rate, "G8")

# denied $ by reason bar
bar = BarChart()
bar.type = "bar"
bar.add_data(Reference(ws_r, min_col=3, min_row=4, max_row=r_last), titles_from_data=True)
bar.set_categories(Reference(ws_r, min_col=1, min_row=5, max_row=r_last))
style_chart(bar, "Denied billed $ by reason")
bar.x_axis.scaling.orientation = "maxMin"   # largest reason on top
bar.y_axis.crosses = "max"                   # keep the value axis at the bottom after the flip
bar.y_axis.number_format = '"$"#,##0,,"M"'
bar.series[0].graphicalProperties.solidFill = BLUE
bar.gapWidth = 60
ws_d.add_chart(bar, "A24")

# plan x network clustered column
col_chart = BarChart()
col_chart.type = "col"
col_chart.add_data(Reference(ws_p, min_col=2, max_col=3, min_row=4, max_row=7), titles_from_data=True)
col_chart.set_categories(Reference(ws_p, min_col=1, min_row=5, max_row=7))
style_chart(col_chart, "Denial rate by plan and network")
col_chart.legend = Legend()
col_chart.legend.position = "t"
col_chart.y_axis.number_format = "0%"
col_chart.series[0].graphicalProperties.solidFill = BLUE
col_chart.series[1].graphicalProperties.solidFill = ORANGE
col_chart.dataLabels = DataLabelList()
col_chart.dataLabels.showVal = True
col_chart.dataLabels.showSerName = col_chart.dataLabels.showCatName = False
col_chart.dataLabels.showLegendKey = False
col_chart.dataLabels.numFmt = "0%"
col_chart.gapWidth = 80
ws_d.add_chart(col_chart, "G24")

# ---------------------------------------------------------------- About
ws_a = wb.create_sheet("About")
title(ws_a, "About this workbook", "Portfolio project: Healthcare Claims Analytics (synthetic data)")
rows = [
    ("Sheet", "What it shows"),
    ("Dashboard", "KPI tiles and four charts, all driven by formulas"),
    ("Monthly", "Claims, denials, denial rate, billed, paid and paid YTD by service month"),
    ("Denial Reasons", "Denied claims and dollars by reason, with 2024 vs 2025 change"),
    ("Plan x Network", "Denial rate matrix: plan type x network status"),
    ("Provider Scorecard", "Every provider vs its peer group; sort by Excess (pts) to find outliers"),
    ("Claims", f"{N:,} cleaned claims (Excel table tblClaims) produced by Python/01_data_cleaning.ipynb"),
    ("", ""),
    ("Definition", ""),
    ("Denial rate", "Denied claims / decided claims (Paid + Denied). Pending claims are excluded."),
    ("Denied billed $", "Billed amount on denied claims: revenue at risk, before appeals"),
    ("Peer rate", "Denial rate of all providers with the same claim type and network status"),
    ("Excess (pts)", "Provider denial rate minus peer rate, in percentage points"),
]
for i, (a, b) in enumerate(rows):
    ws_a[f"A{4 + i}"], ws_a[f"B{4 + i}"] = a, b
    if a in ("Sheet", "Definition"):
        ws_a[f"A{4 + i}"].font = ws_a[f"B{4 + i}"].font = Font(bold=True)
widths(ws_a, {"A": 22, "B": 95})

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)}  ({N:,} claim rows, {len(prov)} providers)")
